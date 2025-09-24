import csv
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

s_classeurIni = r'C:\Users\vt238770\Documents\_CEA\EXCEL - Tests ou Backup\Stats-Eval&Go - Copie.xlsx'
s_classeurDestination = r'C:\Users\vt238770\Documents\_CEA\EXCEL - Tests ou Backup\Stats-Eval&Go - Copie.xlsx'
s_ws = r'Interventions'
s_tableauStructure = s_ws
s_csv = r'P:\FORMATIONS_C\35C\P07-bilan-sessions-et-bilan-formation\rapports-sessions-evaluations\2021-12-S-12664 UEM\S-12664-FC21-35C-VTE-SNA-Interventions.csv'




def colnum_string(n):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string



wb = load_workbook(filename = s_classeurDestination)
ws = wb[s_ws]
print(enumerate(ws._tables))
for i, table in enumerate(ws._tables):
    print(i)
    print(table)
    if table == s_tableauStructure:
        tableRef = i





with open(s_csv, newline='', encoding='latin-1') as f:
    reader = csv.reader(f, delimiter=';')
    for i, row in enumerate(reader):
        print(row)
        for j, cell in enumerate(row): 
            if not i == 0:
                ws.cell(row=i+1, column=j+1).value = cell #float(cell)
            else:
                ws.cell(row=i+1, column=j+1).value = cell

            maxRef = [i,j]

#maxRef = [50,50]


    



#style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
#resTable.tableStyleInfo = style
resTable = Table(displayName="DataVTE", ref="A1:{}{}".format(colnum_string(maxRef[0]), maxRef[1]))


ws._tables[tableRef] = resTable

wb.save('output.xlsx')
wb.close()