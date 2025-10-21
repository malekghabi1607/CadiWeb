import time
import csv
import xlwings as xw
import pandas
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

s_classeurIni = r'C:\Users\vt238770\Documents\_CEA\EXCEL - Tests ou Backup\Stats-Eval&Go - Copie.xlsx'
s_classeurDestination = r'C:\Users\vt238770\Documents\_CEA\EXCEL - Tests ou Backup\Stats-Eval&Go - Copie.xlsx'
s_ws = r'Interventions'
s_tableauStructure = s_ws
s_csv = r'P:\FORMATIONS_C\35C\P07-bilan-sessions-et-bilan-formation\rapports-sessions-evaluations\2021-12-S-12664 UEM\S-12664-FC21-35C-VTE-SNA-Interventions.csv'




#On lit le CSV
data_csv=pandas.read_csv(s_csv, sep=';', encoding='latin-1', skiprows=[0])
print(data_csv)

#wb = load_workbook(filename = s_classeurDestination)
#ws = wb[s_ws]
#table = ws.tables[s_tableauStructure]



# Access the data in the table range
#data = ws[table.ref]
#rows_list = []



#On n'agrandit pas le tableau existant ni les mises en forme conditionnelles + ça crée un fichier excel corrompu
#with pandas.ExcelWriter(s_classeurDestination, mode='a', if_sheet_exists='overlay') as writer:  
#    data_csv.to_excel(writer, sheet_name=s_ws, startrow=1, index=False, header=False)
    
options2 = {}
options2['strings_to_formulas'] = False
options2['strings_to_urls'] = False

writer = pandas.ExcelWriter(s_classeurDestination, mode='a', if_sheet_exists='overlay', engine_kwargs={'keep_vba': True})
data_csv.to_excel(writer, sheet_name=s_ws, startrow=1, index=False, header=False)
writer.close()


#wb.save('output.xlsx')
#wb.close()